"""
Reporte fiscal mensual — totales agrupados por mes para el contador.
Columnas: facturación bruta, comisiones MeLi, costo de envíos,
gastos operativos, costo estimado de fabricación, ganancia neta estimada.

No incluye Ads por mes (requeriría 12 llamadas a la API de MeLi; usá
la pantalla de Publicidad para ese desglose). Todo lo demás viene de
la base de datos local, sin tocar APIs externas.

Los números de ganancia son estimados: el costo de fabricación se
calcula como unidades_vendidas × precio_costo del producto al momento
del reporte (no al momento de la venta).
"""
from datetime import date, timedelta
from io import BytesIO
import db
from utils import formatear_moneda

DIAS_MES_REFERENCIA = 30


def _dias_solapamiento(fecha_inicio, fecha_fin, desde, hasta):
    inicio = max(fecha_inicio, desde)
    fin = min(fecha_fin or hasta, hasta)
    return max((fin - inicio).days + 1, 0)


def calcular_meses(anio: int):
    meses = []
    for mes in range(1, 13):
        desde = date(anio, mes, 1)
        if mes == 12:
            hasta = date(anio, 12, 31)
        else:
            hasta = date(anio, mes + 1, 1) - timedelta(days=1)
        meses.append((desde, hasta))
    return meses


_NOMBRES_MES = ["Enero","Febrero","Marzo","Abril","Mayo","Junio",
                "Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]


def calcular_reporte_anual(usuario_id, anio: int, cuenta_id=None):
    meses = calcular_meses(anio)
    desde_anio = date(anio, 1, 1)
    hasta_anio = date(anio, 12, 31)

    with db.conexion_usuario(usuario_id, cuenta_id) as conexion:
        cursor = conexion.cursor()

        # Costos de fabricación por producto
        cursor.execute("SELECT id_meli, COALESCE(precio_costo, 0) FROM productos_padre")
        costos_fab = dict(cursor.fetchall())

        # Ventas del año — agrupadas por mes
        cursor.execute("""
            SELECT
                EXTRACT(MONTH FROM fecha_venta)::int             AS mes,
                COALESCE(SUM(precio_venta * cantidad), 0)        AS facturacion,
                COALESCE(SUM(cargo_venta), 0)                    AS comisiones,
                COALESCE(SUM(costo_envio), 0)                    AS envios,
                COALESCE(SUM(cantidad), 0)                       AS unidades,
                COUNT(DISTINCT id_orden)                         AS ordenes,
                COALESCE(SUM(retenciones), 0)                    AS retenciones
            FROM ventas
            WHERE fecha_venta BETWEEN %s AND %s
            GROUP BY EXTRACT(MONTH FROM fecha_venta)
        """, (desde_anio, hasta_anio))
        ventas_por_mes = {row[0]: row for row in cursor.fetchall()}

        # Costo fabricación por mes (unidades vendidas × precio_costo por item)
        cursor.execute("""
            SELECT
                EXTRACT(MONTH FROM fecha_venta)::int AS mes,
                id_meli,
                SUM(cantidad)::int                   AS total_uds
            FROM ventas
            WHERE fecha_venta BETWEEN %s AND %s
            GROUP BY EXTRACT(MONTH FROM fecha_venta), id_meli
        """, (desde_anio, hasta_anio))
        fab_por_mes = {}
        for row in cursor.fetchall():
            mes_num, id_meli, total_uds = row[0], row[1], row[2]
            fab_por_mes.setdefault(mes_num, 0)
            fab_por_mes[mes_num] += costos_fab.get(id_meli, 0) * total_uds

        # Gastos operativos del año — todos los que solapan con el año
        cursor.execute("""
            SELECT monto, fecha, recurrente, fecha_fin
            FROM gastos_operativos
            WHERE fecha <= %s AND (fecha_fin IS NULL OR fecha_fin >= %s)
              AND (eliminado_en IS NULL)
        """, (hasta_anio, desde_anio))
        gastos_raw = cursor.fetchall()

    # Calcular gastos por mes en Python (misma lógica que costos.py)
    gastos_por_mes = {}
    for monto_raw, fecha_inicio, recurrente, fecha_fin in gastos_raw:
        monto_base = float(monto_raw or 0)
        if not fecha_inicio:
            continue
        # Normalizar a date
        if hasattr(fecha_inicio, 'date'):
            fecha_inicio = fecha_inicio.date()
        if fecha_fin and hasattr(fecha_fin, 'date'):
            fecha_fin = fecha_fin.date()

        for mes_num, (desde, hasta) in enumerate(meses, 1):
            if recurrente:
                dias = _dias_solapamiento(fecha_inicio, fecha_fin, desde, hasta)
                monto_mes = round(monto_base * (dias / DIAS_MES_REFERENCIA), 2)
            else:
                # Gasto único: aplica al mes en que cae
                if desde <= fecha_inicio <= hasta:
                    monto_mes = monto_base
                else:
                    monto_mes = 0
            gastos_por_mes[mes_num] = gastos_por_mes.get(mes_num, 0) + monto_mes

    resultado = []
    for idx, (desde, hasta) in enumerate(meses):
        mes_num = idx + 1
        fila = ventas_por_mes.get(mes_num)
        facturacion = float(fila[1]) if fila else 0.0
        comisiones  = float(fila[2]) if fila else 0.0
        envios      = float(fila[3]) if fila else 0.0
        unidades    = int(fila[4])   if fila else 0
        ordenes     = int(fila[5])   if fila else 0
        retenciones = float(fila[6]) if fila else 0.0

        costo_fab_total = float(fab_por_mes.get(mes_num, 0))
        gastos          = float(gastos_por_mes.get(mes_num, 0))
        cargos_totales  = comisiones + envios
        ganancia_estimada = facturacion - cargos_totales - gastos - costo_fab_total

        resultado.append({
            "mes": mes_num,
            "nombre": _NOMBRES_MES[idx],
            "desde": desde.isoformat(),
            "hasta": hasta.isoformat(),
            "facturacion": facturacion,
            "comisiones": comisiones,
            "envios": envios,
            "cargos_totales": cargos_totales,
            "gastos": gastos,
            "costo_fabricacion": costo_fab_total,
            "ordenes": ordenes,
            "unidades": unidades,
            "retenciones": retenciones,
            "ganancia_estimada": ganancia_estimada,
            "facturacion_f": formatear_moneda(facturacion),
            "comisiones_f":  formatear_moneda(comisiones),
            "envios_f":      formatear_moneda(envios),
            "gastos_f":      formatear_moneda(gastos),
            "costo_fab_f":   formatear_moneda(costo_fab_total),
            "ganancia_f":    formatear_moneda(ganancia_estimada),
            "retenciones_f": formatear_moneda(retenciones),
        })

    return resultado


def generar_excel_fiscal(meses_data, anio):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = f"Fiscal {anio}"

    encabezado_fill = PatternFill("solid", fgColor="0D0F26")
    encabezado_font = Font(bold=True, color="B06BFF")

    columnas = [
        "Mes", "Órdenes", "Unidades",
        "Facturación Bruta",
        "Comisiones MeLi", "Costo Envíos", "Total Cargos MeLi",
        "Gastos Operativos", "Costo Fabricación (est.)",
        "Ganancia Neta Estimada",
        "Retenciones de impuestos (informativo)",
    ]
    ws.append(columnas)
    for cell in ws[1]:
        cell.font = encabezado_font
        cell.fill = encabezado_fill
        cell.alignment = Alignment(horizontal="center")

    totales = {k: 0.0 for k in ("facturacion","comisiones","envios","cargos_totales","gastos","costo_fabricacion","ganancia_estimada","retenciones")}
    total_ordenes = total_unidades = 0

    for m in meses_data:
        ws.append([
            m["nombre"], m["ordenes"], m["unidades"],
            m["facturacion"], m["comisiones"], m["envios"], m["cargos_totales"],
            m["gastos"], m["costo_fabricacion"], m["ganancia_estimada"], m["retenciones"],
        ])
        for k in totales:
            totales[k] += m[k]
        total_ordenes  += m["ordenes"]
        total_unidades += m["unidades"]

    ws.append([
        "TOTAL AÑO", total_ordenes, total_unidades,
        totales["facturacion"], totales["comisiones"], totales["envios"], totales["cargos_totales"],
        totales["gastos"], totales["costo_fabricacion"], totales["ganancia_estimada"], totales["retenciones"],
    ])
    fila_total = ws.max_row
    total_font = Font(bold=True, color="29E6B0")
    for cell in ws[fila_total]:
        cell.font = total_font

    fmt_moneda = '#,##0.00'
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=4, max_col=11):
        for cell in row:
            cell.number_format = fmt_moneda

    anchos = [14, 9, 10, 20, 18, 15, 18, 20, 24, 24, 34]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho

    ws.append([])
    ws.append(["Las retenciones (IIBB, SIRTAC y similares) las descuenta Mercado Libre al acreditar: no se restan de la ganancia porque son pago anticipado de impuestos. Confirmá su tratamiento con tu contador."])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def detalle_del_mes(usuario_id, anio: int, mes: int, cuenta_id=None):
    """Una fila por línea de venta del mes (el "libro de ventas" que suele pedir el contador). Sin datos personales de compradores."""
    desde = date(anio, mes, 1)
    hasta = (date(anio + 1, 1, 1) if mes == 12 else date(anio, mes + 1, 1)) - timedelta(days=1)
    with db.conexion_usuario(usuario_id, cuenta_id) as conexion:
        cursor = conexion.cursor()
        cursor.execute("""
            SELECT fecha_venta, id_orden, titulo, cantidad, precio_venta, cargo_venta, costo_envio, retenciones, neto_recibido, origen
            FROM ventas WHERE fecha_venta BETWEEN %s AND %s ORDER BY fecha_venta, id_orden, id
        """, (desde, hasta))
        filas = cursor.fetchall()
    return [{"fecha": f, "orden": o, "producto": t, "cantidad": int(c or 0), "importe": float(p or 0) * int(c or 0), "cargo_meli": float(cv or 0),
             "envio": float(ce or 0), "retenciones": float(r or 0), "neto_recibido": float(n) if n is not None else None, "origen": og}
            for f, o, t, c, p, cv, ce, r, n, og in filas]


def generar_excel_detalle(filas, anio: int, mes: int):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = f"Ventas {_NOMBRES_MES[mes - 1]} {anio}"[:31]
    ws.append(["Fecha", "N.º de orden", "Producto", "Cantidad", "Importe de la venta", "Cargo de Mercado Libre", "Envío", "Retenciones", "Neto acreditado", "Origen"])
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center")
    for f in filas:
        ws.append([f["fecha"], f["orden"], f["producto"], f["cantidad"], f["importe"], f["cargo_meli"], f["envio"], f["retenciones"],
                   f["neto_recibido"], "Mercado Libre" if f["origen"] == "meli" else "Manual"])
    ultima = ws.max_row
    ws.append(["TOTAL", "", "", sum(f["cantidad"] for f in filas), sum(f["importe"] for f in filas), sum(f["cargo_meli"] for f in filas),
               sum(f["envio"] for f in filas), sum(f["retenciones"] for f in filas), sum(f["neto_recibido"] or 0 for f in filas), ""])
    for celda in ws[ws.max_row]:
        celda.font = Font(bold=True)
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=5, max_col=9):
        for celda in fila:
            celda.number_format = "#,##0.00"
    for fila in ws.iter_rows(min_row=2, max_row=ultima, min_col=1, max_col=1):
        fila[0].number_format = "dd/mm/yyyy"
    for i, ancho in enumerate([12, 20, 52, 10, 20, 22, 14, 14, 16, 14], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
