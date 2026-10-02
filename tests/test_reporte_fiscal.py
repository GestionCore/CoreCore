import io
from datetime import date

from openpyxl import load_workbook

import reporte_fiscal as rf
from tests.conftest import CursorFalso


def test_el_excel_de_detalle_trae_una_fila_por_venta_y_los_totales():
    filas = [
        {"fecha": date(2026, 9, 3), "orden": "2000001", "producto": "Campera", "cantidad": 2, "importe": 40000.0, "cargo_meli": 7000.0,
         "envio": 3000.0, "retenciones": 600.0, "neto_recibido": 29400.0, "origen": "meli"},
        {"fecha": date(2026, 9, 4), "orden": "M-1", "producto": "Remera", "cantidad": 1, "importe": 10000.0, "cargo_meli": 0.0,
         "envio": 0.0, "retenciones": 0.0, "neto_recibido": None, "origen": "manual"},
    ]
    ws = load_workbook(io.BytesIO(rf.generar_excel_detalle(filas, 2026, 9).read())).active
    datos = list(ws.iter_rows(values_only=True))
    assert datos[0][:3] == ("Fecha", "N.º de orden", "Producto")
    assert datos[1][9] == "Mercado Libre" and datos[2][9] == "Manual"
    assert datos[-1][0] == "TOTAL" and datos[-1][3] == 3 and datos[-1][4] == 50000.0 and datos[-1][8] == 29400.0
    assert "Septiembre" in ws.title


def test_el_detalle_no_incluye_datos_personales_del_comprador():
    consulta = []

    class Cur(CursorFalso):
        def execute(self, sql, params=None):
            consulta.append(sql)
            super().execute(sql, params)

    class Con:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def cursor(self):
            return Cur([])

    import db
    original = db.conexion_usuario
    db.conexion_usuario = lambda *a, **k: Con()
    try:
        assert rf.detalle_del_mes(1, 2026, 12, 1) == []
    finally:
        db.conexion_usuario = original
    sql = consulta[0].lower()
    assert "comprador" not in sql and "provincia" not in sql
